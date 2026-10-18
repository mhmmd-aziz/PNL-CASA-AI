import io
import json
import sqlite3
import sys
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

import auth_routes
import dashboard_routes as dr

SCHEMA = """
CREATE TABLE examinations (id INTEGER PRIMARY KEY, sample_code TEXT, status TEXT, created_at TEXT);
CREATE TABLE analysis_runs (id INTEGER PRIMARY KEY, examination_id INTEGER, status TEXT, completed_at TEXT);
CREATE TABLE reports (id INTEGER PRIMARY KEY, examination_id INTEGER, released_at TEXT);
CREATE TABLE audit_logs (id INTEGER PRIMARY KEY, action_type TEXT, result TEXT, created_at TEXT);
"""


def make_db():
    conn = sqlite3.connect(":memory:", check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.executescript(SCHEMA)
    conn.executemany(
        "INSERT INTO examinations (id, sample_code, status, created_at) VALUES (?, ?, ?, ?)",
        [
            # 18:30 UTC = 01:30 WIB hari berikutnya
            (1, "SMP-A", "RELEASED", "2026-10-01 18:30:00"),
            (2, "SMP-B", "CANCELLED", "2026-10-02 03:00:00"),
            (3, "SMP-C", "PENDING_REVIEW", "2026-10-03 03:00:00"),
            (4, "SMP-OLD", "RELEASED", "2026-08-15 03:00:00"),
        ],
    )
    conn.executemany(
        "INSERT INTO analysis_runs (examination_id, status, completed_at) VALUES (?, ?, ?)",
        [
            (1, "COMPLETED", "2026-10-02T02:00:00+00:00"),
            (3, "COMPLETED", "2026-10-03T04:00:00+00:00"),
            (3, "FAILED", "2026-10-03T05:00:00+00:00"),
        ],
    )
    # Sampel A: registrasi 18:30 UTC, rilis 10 jam kemudian
    conn.execute("INSERT INTO reports (examination_id, released_at) VALUES (1, '2026-10-02T04:30:00+00:00')")
    conn.executemany(
        "INSERT INTO audit_logs (action_type, result, created_at) VALUES (?, ?, ?)",
        [
            ("LOGIN_FAILED", "FAILURE", "2026-10-02 01:00:00"),
            ("LOGIN_FAILED", "FAILURE", "2026-10-02 02:00:00"),
            ("LOGIN_SUCCESS", "SUCCESS", "2026-10-02 02:05:00"),
            ("OVERRIDE_SAME_DAY_BOOKING", "SUCCESS", "2026-10-03 02:00:00"),
            ("LOGIN_FAILED", "FAILURE", "2026-09-01 02:00:00"),
        ],
    )
    conn.commit()
    return conn


class BuildDashboardTests(unittest.TestCase):
    def setUp(self):
        self.conn = make_db()

    def test_daily_series_uses_wib_dates_and_fills_empty_days(self):
        data = dr.build_dashboard(self.conn, date(2026, 10, 1), date(2026, 10, 4))
        series = {item["bucket"]: item for item in data["series"]}
        self.assertEqual(list(series), ["2026-10-01", "2026-10-02", "2026-10-03", "2026-10-04"])
        self.assertEqual(series["2026-10-01"]["registered"], 0)  # 18:30 UTC tanggal 1 jatuh ke 2 WIB
        self.assertEqual(series["2026-10-02"]["registered"], 2)
        self.assertEqual(series["2026-10-03"]["registered"], 1)
        self.assertEqual(series["2026-10-04"], {"bucket": "2026-10-04", "registered": 0, "analyzed": 0, "released": 0})
        self.assertEqual(data["range"]["granularity"], "day")

    def test_summary_counts_only_completed_analyses_and_released_reports(self):
        summary = dr.build_dashboard(self.conn, date(2026, 10, 1), date(2026, 10, 31))["summary"]
        self.assertEqual(summary["registered"], 3)
        self.assertEqual(summary["analyzed"], 2)
        self.assertEqual(summary["released"], 1)
        self.assertEqual(summary["avg_turnaround_hours"], 10.0)
        self.assertEqual(summary["turnaround_samples"], 1)

    def test_status_filter_limits_every_examination_metric(self):
        data = dr.build_dashboard(self.conn, date(2026, 10, 1), date(2026, 10, 31), status="PENDING_REVIEW")
        self.assertEqual(data["summary"]["registered"], 1)
        self.assertEqual(data["summary"]["analyzed"], 1)
        self.assertEqual(data["summary"]["released"], 0)
        self.assertEqual([item["sample_code"] for item in data["examinations"]], ["SMP-C"])
        self.assertEqual(data["status_distribution"], [{"status": "PENDING_REVIEW", "label": "Menunggu tinjauan dokter", "count": 1}])

    def test_security_counts_follow_date_range(self):
        security = dr.build_dashboard(self.conn, date(2026, 10, 1), date(2026, 10, 31))["security"]
        self.assertEqual(security["login_failed"], 2)
        self.assertEqual(security["overrides"], 1)
        self.assertEqual(security["failed_actions"], 2)
        self.assertEqual(security["audit_events"], 4)
        self.assertEqual(security["top_actions"][0], {"action_type": "LOGIN_FAILED", "count": 2, "failures": 2})

    def test_long_range_groups_by_month(self):
        data = dr.build_dashboard(self.conn, date(2026, 8, 1), date(2026, 12, 31))
        self.assertEqual(data["range"]["granularity"], "month")
        self.assertEqual([item["bucket"] for item in data["series"]], ["2026-08", "2026-09", "2026-10", "2026-11", "2026-12"])
        self.assertEqual(data["series"][0]["registered"], 1)
        self.assertEqual(data["series"][2]["registered"], 3)

    def test_month_buckets_cross_year_boundary(self):
        self.assertEqual(dr._buckets(date(2026, 11, 20), date(2027, 2, 1), "month"), ["2026-11", "2026-12", "2027-01", "2027-02"])

    def test_invalid_ranges_are_rejected(self):
        with self.assertRaises(HTTPException) as ctx:
            dr.build_dashboard(self.conn, date(2026, 10, 5), date(2026, 10, 1))
        self.assertEqual(ctx.exception.status_code, 422)
        with self.assertRaises(HTTPException):
            dr.build_dashboard(self.conn, date(2020, 1, 1), date(2026, 1, 1))

    def test_single_day_range_is_allowed(self):
        data = dr.build_dashboard(self.conn, date(2026, 10, 2), date(2026, 10, 2))
        self.assertEqual(len(data["series"]), 1)

    def test_examination_list_is_truncated_and_flagged(self):
        data = dr.build_dashboard(self.conn, date(2026, 10, 1), date(2026, 10, 31), list_limit=2)
        self.assertEqual(len(data["examinations"]), 2)
        self.assertTrue(data["examinations_truncated"])
        self.assertEqual(data["examinations"][0]["sample_code"], "SMP-C")  # terbaru dahulu

    def test_empty_range_returns_zeroes_not_errors(self):
        data = dr.build_dashboard(self.conn, date(2025, 1, 1), date(2025, 1, 7))
        self.assertEqual(data["summary"]["registered"], 0)
        self.assertIsNone(data["summary"]["avg_turnaround_hours"])
        self.assertEqual(data["status_distribution"], [])
        self.assertEqual(data["examinations"], [])

    def test_output_contains_no_patient_identity_fields(self):
        text = json.dumps(dr.build_dashboard(self.conn, date(2026, 10, 1), date(2026, 10, 31)))
        for field in ("full_name", "medical_record_number", "date_of_birth", "patient"):
            self.assertNotIn(field, text)


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.data = dr.build_dashboard(make_db(), date(2026, 10, 1), date(2026, 10, 31))

    def test_xlsx_has_expected_sheets_and_values(self):
        workbook = load_workbook(io.BytesIO(dr.export_xlsx(self.data)))
        self.assertEqual(workbook.sheetnames, ["Ringkasan", "Tren", "Status", "Audit", "Pemeriksaan"])
        summary = {row[0]: row[1] for row in workbook["Ringkasan"].iter_rows(values_only=True) if row and row[0]}
        self.assertEqual(summary["Pemeriksaan terdaftar"], 3)
        self.assertEqual(summary["Laporan dirilis"], 1)
        trend = list(workbook["Tren"].iter_rows(values_only=True))
        self.assertEqual(trend[0], ("Tanggal", "Pemeriksaan terdaftar", "Analisis selesai", "Laporan dirilis"))
        self.assertEqual(len(trend), 1 + 31)
        codes = [row[0] for row in workbook["Pemeriksaan"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(sorted(codes), ["SMP-A", "SMP-B", "SMP-C"])

    def test_nice_axis_maximum_always_covers_the_peak(self):
        for peak, expected in ((0, 4), (3, 4), (7, 10), (15, 20), (23, 25), (180, 200), (1200, 2000)):
            with self.subTest(peak=peak):
                self.assertEqual(dr._nice_max(peak), expected)
                self.assertGreaterEqual(dr._nice_max(peak), peak)

    def test_pdf_is_valid_document(self):
        content = dr.export_pdf(self.data)
        self.assertTrue(content.startswith(b"%PDF"))
        self.assertGreater(len(content), 1500)

    def test_exports_work_for_empty_and_monthly_ranges(self):
        for start, end in ((date(2025, 1, 1), date(2025, 1, 3)), (date(2026, 1, 1), date(2026, 12, 31))):
            with self.subTest(start=start, end=end):
                data = dr.build_dashboard(make_db(), start, end)
                self.assertTrue(dr.export_pdf(data).startswith(b"%PDF"))
                load_workbook(io.BytesIO(dr.export_xlsx(data)))


class EndpointTests(unittest.TestCase):
    def setUp(self):
        self.conn = make_db()
        app = FastAPI()
        app.include_router(dr.router)
        self.app = app
        self.client = TestClient(app)
        self.role = "SUPER_ADMIN"
        app.dependency_overrides[auth_routes.get_current_user] = lambda: {
            "id": 1, "email": "admin@lab.local", "role": self.role, "full_name": "Admin",
        }
        self.patches = [
            patch.object(dr, "get_connection", return_value=self.conn),
            patch.object(dr, "log_audit_event"),
        ]
        self.audit = None
        for item in self.patches:
            mock = item.start()
            if item.attribute == "log_audit_event":
                self.audit = mock

    def tearDown(self):
        for item in self.patches:
            item.stop()

    def test_super_admin_gets_dashboard(self):
        response = self.client.get("/api/admin/dashboard", params={"start": "2026-10-01", "end": "2026-10-31"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["summary"]["registered"], 3)

    def test_other_roles_are_denied(self):
        for role in ("FRONT_OFFICE", "ANALYST", "PHYSICIAN", "PATIENT"):
            with self.subTest(role=role):
                self.role = role
                response = self.client.get("/api/admin/dashboard", params={"start": "2026-10-01", "end": "2026-10-31"})
                self.assertEqual(response.status_code, 403)
                export = self.client.get("/api/admin/dashboard/export", params={"format": "pdf", "start": "2026-10-01", "end": "2026-10-31"})
                self.assertEqual(export.status_code, 403)
        self.audit.assert_not_called()

    def test_invalid_status_and_dates_are_rejected(self):
        bad_status = self.client.get("/api/admin/dashboard", params={"start": "2026-10-01", "end": "2026-10-31", "status": "DROP"})
        self.assertEqual(bad_status.status_code, 422)
        reversed_range = self.client.get("/api/admin/dashboard", params={"start": "2026-10-31", "end": "2026-10-01"})
        self.assertEqual(reversed_range.status_code, 422)
        missing = self.client.get("/api/admin/dashboard", params={"start": "2026-10-01"})
        self.assertEqual(missing.status_code, 422)

    def test_export_returns_files_and_writes_audit(self):
        for fmt, media in (("pdf", "application/pdf"), ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")):
            with self.subTest(fmt=fmt):
                response = self.client.get(
                    "/api/admin/dashboard/export",
                    params={"format": fmt, "start": "2026-10-01", "end": "2026-10-31", "status": "RELEASED"},
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.headers["content-type"], media)
                self.assertIn(f'laporan-casa-ai_2026-10-01_2026-10-31.{fmt}', response.headers["content-disposition"])
        self.assertEqual(self.audit.call_count, 2)
        self.assertEqual(self.audit.call_args.kwargs["action_type"], "EXPORT_DASHBOARD_REPORT")

    def test_unknown_export_format_is_rejected(self):
        response = self.client.get("/api/admin/dashboard/export", params={"format": "csv", "start": "2026-10-01", "end": "2026-10-31"})
        self.assertEqual(response.status_code, 422)


if __name__ == "__main__":
    unittest.main()
