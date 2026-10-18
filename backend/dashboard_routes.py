"""
Dashboard Super Admin: ringkasan operasional laboratorium per rentang tanggal,
beserta unduhan laporan PDF dan Excel dengan filter yang sama.

Isi bersifat operasional (volume, status, waktu proses, kejadian keamanan).
Laporan tidak memuat nama atau nomor rekam medis pasien; sampel hanya dirujuk
dengan kode sampel. Tanggal dihitung dalam WIB (UTC+7) karena data disimpan UTC.
"""
import io
from datetime import date, datetime, timedelta, timezone
from typing import Literal, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response

from auth_routes import require_roles
from database.db_connection import get_connection, log_audit_event

router = APIRouter(prefix="/api/admin/dashboard", tags=["admin-dashboard"])

SUPER_ADMIN = Depends(require_roles("SUPER_ADMIN"))
WIB = timezone(timedelta(hours=7))
WIB_MODIFIER = "+7 hours"
MAX_RANGE_DAYS = 366 * 5
DAILY_MAX_DAYS = 92  # rentang lebih panjang dikelompokkan per bulan
LIST_LIMIT = 200
EXPORT_LIST_LIMIT = 5000

ExamStatus = Literal[
    "REGISTERED", "IN_PROGRESS", "PENDING_REVIEW", "RETURNED_FOR_CORRECTION",
    "APPROVED", "RELEASED", "CANCELLED", "FAILED_ANALYSIS",
]
STATUS_LABELS = {
    "REGISTERED": "Terdaftar",
    "IN_PROGRESS": "Sedang diperiksa",
    "PENDING_REVIEW": "Menunggu tinjauan dokter",
    "RETURNED_FOR_CORRECTION": "Dikembalikan untuk koreksi",
    "APPROVED": "Disetujui dokter",
    "RELEASED": "Laporan dirilis",
    "CANCELLED": "Dibatalkan",
    "FAILED_ANALYSIS": "Analisis gagal",
}


def validate_range(start: date, end: date) -> None:
    if end < start:
        raise HTTPException(status_code=422, detail="Tanggal akhir tidak boleh sebelum tanggal awal.")
    if (end - start).days + 1 > MAX_RANGE_DAYS:
        raise HTTPException(status_code=422, detail="Rentang tanggal maksimal 5 tahun.")


def _granularity(start: date, end: date) -> str:
    return "day" if (end - start).days + 1 <= DAILY_MAX_DAYS else "month"


def _buckets(start: date, end: date, granularity: str) -> list:
    if granularity == "day":
        return [(start + timedelta(days=i)).isoformat() for i in range((end - start).days + 1)]
    buckets, year, month = [], start.year, start.month
    while (year, month) <= (end.year, end.month):
        buckets.append(f"{year:04d}-{month:02d}")
        year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    return buckets


def _bucket_sql(column: str, granularity: str) -> str:
    if granularity == "day":
        return f"date({column}, '{WIB_MODIFIER}')"
    return f"strftime('%Y-%m', {column}, '{WIB_MODIFIER}')"


def _in_range_sql(column: str) -> str:
    return f"date({column}, '{WIB_MODIFIER}') BETWEEN ? AND ?"


def _counts(conn, sql: str, params: tuple) -> dict:
    return {row[0]: row[1] for row in conn.execute(sql, params).fetchall() if row[0] is not None}


def build_dashboard(conn, start: date, end: date, status: Optional[str] = None, list_limit: int = LIST_LIMIT) -> dict:
    validate_range(start, end)
    granularity = _granularity(start, end)
    span = (start.isoformat(), end.isoformat())
    status_sql = " AND e.status = ?" if status else ""
    status_params = (status,) if status else ()

    registered = _counts(
        conn,
        f"SELECT {_bucket_sql('e.created_at', granularity)}, COUNT(*) FROM examinations e "
        f"WHERE {_in_range_sql('e.created_at')}{status_sql} GROUP BY 1",
        span + status_params,
    )
    analyzed = _counts(
        conn,
        f"SELECT {_bucket_sql('ar.completed_at', granularity)}, COUNT(*) FROM analysis_runs ar "
        f"JOIN examinations e ON e.id = ar.examination_id "
        f"WHERE ar.status = 'COMPLETED' AND ar.completed_at IS NOT NULL AND {_in_range_sql('ar.completed_at')}{status_sql} "
        f"GROUP BY 1",
        span + status_params,
    )
    released = _counts(
        conn,
        f"SELECT {_bucket_sql('r.released_at', granularity)}, COUNT(*) FROM reports r "
        f"JOIN examinations e ON e.id = r.examination_id "
        f"WHERE r.released_at IS NOT NULL AND {_in_range_sql('r.released_at')}{status_sql} GROUP BY 1",
        span + status_params,
    )
    series = [
        {
            "bucket": bucket,
            "registered": registered.get(bucket, 0),
            "analyzed": analyzed.get(bucket, 0),
            "released": released.get(bucket, 0),
        }
        for bucket in _buckets(start, end, granularity)
    ]

    status_rows = conn.execute(
        f"SELECT e.status, COUNT(*) FROM examinations e WHERE {_in_range_sql('e.created_at')}{status_sql} "
        f"GROUP BY e.status ORDER BY COUNT(*) DESC, e.status",
        span + status_params,
    ).fetchall()
    status_distribution = [
        {"status": row[0], "label": STATUS_LABELS.get(row[0], row[0]), "count": row[1]} for row in status_rows
    ]

    turnaround = conn.execute(
        f"SELECT AVG((julianday(r.released_at) - julianday(e.created_at)) * 24), COUNT(*) FROM reports r "
        f"JOIN examinations e ON e.id = r.examination_id "
        f"WHERE r.released_at IS NOT NULL AND {_in_range_sql('r.released_at')}{status_sql}",
        span + status_params,
    ).fetchone()

    audit_span = f"date(created_at, '{WIB_MODIFIER}') BETWEEN ? AND ?"
    security = conn.execute(
        f"""
        SELECT
            COALESCE(SUM(CASE WHEN action_type = 'LOGIN_FAILED' THEN 1 ELSE 0 END), 0),
            COALESCE(SUM(CASE WHEN action_type LIKE 'OVERRIDE_%' THEN 1 ELSE 0 END), 0),
            COALESCE(SUM(CASE WHEN result <> 'SUCCESS' THEN 1 ELSE 0 END), 0),
            COUNT(*)
        FROM audit_logs WHERE {audit_span}
        """,
        span,
    ).fetchone()
    top_actions = [
        {"action_type": row[0], "count": row[1], "failures": row[2]}
        for row in conn.execute(
            f"""
            SELECT action_type, COUNT(*), SUM(CASE WHEN result <> 'SUCCESS' THEN 1 ELSE 0 END)
            FROM audit_logs WHERE {audit_span}
            GROUP BY action_type ORDER BY COUNT(*) DESC, action_type LIMIT 10
            """,
            span,
        ).fetchall()
    ]

    exam_rows = conn.execute(
        f"""
        SELECT e.sample_code, e.status, datetime(e.created_at, '{WIB_MODIFIER}'),
               (SELECT datetime(MAX(r.released_at), '{WIB_MODIFIER}') FROM reports r
                WHERE r.examination_id = e.id AND r.released_at IS NOT NULL)
        FROM examinations e
        WHERE {_in_range_sql('e.created_at')}{status_sql}
        ORDER BY e.created_at DESC, e.id DESC LIMIT ?
        """,
        span + status_params + (list_limit + 1,),
    ).fetchall()
    examinations = [
        {
            "sample_code": row[0],
            "status": row[1],
            "status_label": STATUS_LABELS.get(row[1], row[1]),
            "registered_at": row[2],
            "released_at": row[3],
        }
        for row in exam_rows[:list_limit]
    ]

    return {
        "range": {"start": span[0], "end": span[1], "granularity": granularity, "timezone": "WIB (UTC+7)"},
        "filters": {"status": status, "status_label": STATUS_LABELS.get(status) if status else "Semua status"},
        "generated_at": datetime.now(WIB).strftime("%Y-%m-%d %H:%M WIB"),
        "summary": {
            "registered": sum(item["registered"] for item in series),
            "analyzed": sum(item["analyzed"] for item in series),
            "released": sum(item["released"] for item in series),
            "avg_turnaround_hours": round(turnaround[0], 1) if turnaround[0] is not None else None,
            "turnaround_samples": turnaround[1],
        },
        "series": series,
        "status_distribution": status_distribution,
        "security": {
            "login_failed": security[0],
            "overrides": security[1],
            "failed_actions": security[2],
            "audit_events": security[3],
            "top_actions": top_actions,
        },
        "examinations": examinations,
        "examinations_truncated": len(exam_rows) > list_limit,
    }


# ---------------------------------------------------------------------------
# Ekspor
# ---------------------------------------------------------------------------

SERIES_COLUMNS = (("registered", "Pemeriksaan terdaftar"), ("analyzed", "Analisis selesai"), ("released", "Laporan dirilis"))
SERIES_COLORS = ("#2a78d6", "#eb6834", "#1baf7a")


def _summary_rows(data: dict) -> list:
    summary, security = data["summary"], data["security"]
    turnaround = summary["avg_turnaround_hours"]
    return [
        ("Periode", f"{data['range']['start']} s.d. {data['range']['end']} ({data['range']['timezone']})"),
        ("Filter status", data["filters"]["status_label"]),
        ("Dibuat", data["generated_at"]),
        ("Pemeriksaan terdaftar", summary["registered"]),
        ("Analisis selesai", summary["analyzed"]),
        ("Laporan dirilis", summary["released"]),
        ("Rata-rata waktu registrasi s.d. rilis (jam)", "-" if turnaround is None else turnaround),
        ("Login gagal", security["login_failed"]),
        ("Pemaksaan (override)", security["overrides"]),
        ("Aksi gagal tercatat", security["failed_actions"]),
    ]


def _bucket_label(data: dict) -> str:
    return "Tanggal" if data["range"]["granularity"] == "day" else "Bulan"


def export_xlsx(data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    bold = Font(bold=True)
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Ringkasan"
    sheet.append(["Laporan Operasional CASA-AI"])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([])
    for label, value in _summary_rows(data):
        sheet.append([label, value])
    sheet.append([])
    sheet.append(["Catatan: laporan operasional, bukan hasil klinis. Tidak memuat identitas pasien."])
    sheet.column_dimensions["A"].width = 44
    sheet.column_dimensions["B"].width = 44

    def table(title: str, header: list, rows: list, widths: list):
        ws = workbook.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font = bold
        for row in rows:
            ws.append(list(row))
        for index, width in enumerate(widths):
            ws.column_dimensions[chr(ord("A") + index)].width = width
        ws.freeze_panes = "A2"

    table(
        "Tren",
        [_bucket_label(data)] + [label for _, label in SERIES_COLUMNS],
        [[item["bucket"]] + [item[key] for key, _ in SERIES_COLUMNS] for item in data["series"]],
        [14, 22, 18, 18],
    )
    table("Status", ["Status", "Jumlah"], [[item["label"], item["count"]] for item in data["status_distribution"]], [32, 12])
    table(
        "Audit",
        ["Aksi", "Jumlah", "Gagal"],
        [[item["action_type"], item["count"], item["failures"]] for item in data["security"]["top_actions"]],
        [36, 12, 12],
    )
    table(
        "Pemeriksaan",
        ["Kode sampel", "Status", "Terdaftar (WIB)", "Dirilis (WIB)"],
        [[item["sample_code"], item["status_label"], item["registered_at"], item["released_at"] or "-"] for item in data["examinations"]],
        [26, 28, 20, 20],
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _trend_drawing(data: dict):
    from reportlab.graphics.charts.legends import Legend
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors

    series = data["series"]
    drawing = Drawing(480, 235)
    chart = HorizontalLineChart()
    chart.x, chart.y, chart.width, chart.height = 40, 50, 420, 150
    chart.data = [[item[key] for item in series] for key, _ in SERIES_COLUMNS]
    step = max(1, len(series) // 8)
    chart.categoryAxis.categoryNames = [item["bucket"] if i % step == 0 else "" for i, item in enumerate(series)]
    chart.categoryAxis.labels.angle = 30
    chart.categoryAxis.labels.boxAnchor = "ne"
    chart.categoryAxis.labels.fontName = "Helvetica"
    chart.categoryAxis.labels.fontSize = 7
    top = _nice_max(max(max(row) for row in chart.data))
    chart.valueAxis.valueMin = 0
    chart.valueAxis.valueMax = top
    chart.valueAxis.valueStep = top / 4
    chart.valueAxis.labels.fontName = "Helvetica"
    chart.valueAxis.labels.fontSize = 7
    chart.valueAxis.labelTextFormat = lambda value: f"{value:g}"
    chart.valueAxis.visibleGrid = True
    chart.valueAxis.gridStrokeColor = colors.HexColor("#e1e0d9")
    for index, color in enumerate(SERIES_COLORS):
        chart.lines[index].strokeColor = colors.HexColor(color)
        chart.lines[index].strokeWidth = 2
    drawing.add(chart)

    # Legenda di atas plot agar tidak bertumpuk dengan label tanggal yang miring
    legend = Legend()
    legend.x, legend.y = 40, 228
    legend.boxAnchor = "nw"
    legend.alignment = "right"
    legend.fontName = "Helvetica"
    legend.fontSize = 8
    legend.columnMaximum = 1
    legend.deltax = 140
    legend.dxTextSpace = 4
    legend.dx = 10
    legend.dy = 3
    legend.colorNamePairs = [(colors.HexColor(color), label) for color, (_, label) in zip(SERIES_COLORS, SERIES_COLUMNS)]
    drawing.add(legend)
    return drawing


def _nice_max(value: float) -> float:
    """Batas atas sumbu yang bulat (1, 2, 2.5, 5 x 10^n), minimal 4 agar grid tetap terbaca."""
    if value <= 4:
        return 4
    magnitude = 10 ** len(str(int(value))) / 10
    return next(m * magnitude for m in (1, 2, 2.5, 5, 10) if m * magnitude >= value)


def export_pdf(data: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=1.6 * cm, rightMargin=1.6 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        title="Laporan Operasional CASA-AI",
    )
    grid = TableStyle([
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8),
        ("FONT", (0, 1), (-1, -1), "Helvetica", 8),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0efec")),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#e1e0d9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ])

    def section(title: str):
        return [Spacer(1, 10), Paragraph(title, styles["Heading3"])]

    def table(header: list, rows: list, widths: list):
        result = Table([header] + [[str(cell) for cell in row] for row in rows], colWidths=widths, repeatRows=1)
        result.setStyle(grid)
        return result

    story = [
        Paragraph("Laporan Operasional CASA-AI", styles["Title"]),
        Paragraph("Laporan operasional, bukan hasil klinis. Tidak memuat identitas pasien.", styles["Italic"]),
    ]
    summary = Table([[label, str(value)] for label, value in _summary_rows(data)], colWidths=[8 * cm, 9.5 * cm])
    summary.setStyle(TableStyle([
        ("FONT", (0, 0), (0, -1), "Helvetica-Bold", 8),
        ("FONT", (1, 0), (1, -1), "Helvetica", 8),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#e1e0d9")),
    ]))
    story += section("Ringkasan") + [summary]
    story += section("Tren per " + _bucket_label(data).lower()) + [_trend_drawing(data)]
    story.append(table(
        [_bucket_label(data)] + [label for _, label in SERIES_COLUMNS],
        [[item["bucket"]] + [item[key] for key, _ in SERIES_COLUMNS] for item in data["series"]],
        [3.5 * cm, 4.5 * cm, 4.5 * cm, 4.5 * cm],
    ))
    story += section("Distribusi status") + [
        table(["Status", "Jumlah"], [[item["label"], item["count"]] for item in data["status_distribution"]] or [["-", 0]], [10 * cm, 4 * cm])
    ]
    story += section("Aktivitas audit teratas") + [
        table(
            ["Aksi", "Jumlah", "Gagal"],
            [[item["action_type"], item["count"], item["failures"]] for item in data["security"]["top_actions"]] or [["-", 0, 0]],
            [9 * cm, 3 * cm, 3 * cm],
        )
    ]
    story += section("Daftar pemeriksaan") + [
        table(
            ["Kode sampel", "Status", "Terdaftar (WIB)", "Dirilis (WIB)"],
            [[item["sample_code"], item["status_label"], item["registered_at"], item["released_at"] or "-"] for item in data["examinations"]] or [["-", "-", "-", "-"]],
            [4.6 * cm, 4.6 * cm, 4 * cm, 4 * cm],
        )
    ]
    if data["examinations_truncated"]:
        story.append(Paragraph(f"Daftar dipotong pada {len(data['examinations'])} baris pertama.", styles["Italic"]))
    doc.build(story)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Endpoint
# ---------------------------------------------------------------------------

@router.get("")
def dashboard(
    start: date = Query(...),
    end: date = Query(...),
    status: Optional[ExamStatus] = Query(default=None),
    user: dict = SUPER_ADMIN,
):
    return build_dashboard(get_connection(), start, end, status)


@router.get("/export")
def export_dashboard(
    format: Literal["pdf", "xlsx"] = Query(...),
    start: date = Query(...),
    end: date = Query(...),
    status: Optional[ExamStatus] = Query(default=None),
    user: dict = SUPER_ADMIN,
):
    data = build_dashboard(get_connection(), start, end, status, list_limit=EXPORT_LIST_LIMIT)
    if format == "pdf":
        content, media_type = export_pdf(data), "application/pdf"
    else:
        content, media_type = export_xlsx(data), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    file_name = f"laporan-casa-ai_{start.isoformat()}_{end.isoformat()}.{format}"
    log_audit_event(
        actor_id=user["id"],
        action_type="EXPORT_DASHBOARD_REPORT",
        description=f"Laporan dashboard {format.upper()} diunduh untuk {start.isoformat()} s.d. {end.isoformat()} ({data['filters']['status_label']})",
        target_type="dashboard",
        actor_email=user["email"],
        actor_role=user["role"],
    )
    return Response(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{file_name}"'},
    )
